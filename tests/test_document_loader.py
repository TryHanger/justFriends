"""Regression coverage for document extraction and its character limit."""

import asyncio
from io import BytesIO

import pytest
from docx import Document
from fastapi import HTTPException, UploadFile

from app.services.document_loader import extract_text


def test_docx_extracts_paragraphs_and_nested_tables_in_document_order():
    document = Document()
    document.add_paragraph("Before")
    outer = document.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "First cell"
    outer.cell(0, 1).text = "Second cell"
    merged = outer.cell(1, 0).merge(outer.cell(1, 1))
    merged.text = "Merged once"
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell"
    document.add_paragraph("After")
    data = BytesIO()
    document.save(data)

    text, filename = asyncio.run(
        extract_text(UploadFile(filename="poem.docx", file=BytesIO(data.getvalue())), 20)
    )

    assert filename == "poem.docx"
    assert text.splitlines() == [
        "Before", "First cell", "Nested cell", "Second cell", "Merged once", "After"
    ]


def test_extracted_text_limit_counts_trimmed_characters():
    accepted, _ = asyncio.run(
        extract_text(
            UploadFile(filename="limit.txt", file=BytesIO(("  " + "x" * 100_000 + "  ").encode())),
            20,
        )
    )
    assert len(accepted) == 100_000

    with pytest.raises(HTTPException) as error:
        asyncio.run(
            extract_text(
                UploadFile(filename="limit.txt", file=BytesIO(("x" * 100_001).encode())),
                20,
            )
        )
    assert error.value.status_code == 422
    assert "100000" in str(error.value.detail).replace(",", "").replace("_", "")


def test_docx_keeps_internal_empty_body_paragraphs():
    document = Document()
    document.add_paragraph("First")
    document.add_paragraph("")
    document.add_paragraph("Last")
    data = BytesIO()
    document.save(data)

    text, _ = asyncio.run(
        extract_text(UploadFile(filename="poem.docx", file=BytesIO(data.getvalue())), 20)
    )
    assert text == "First\n\nLast"


def test_docx_keeps_every_distinct_table_cell():
    document = Document()
    table = document.add_table(rows=10, cols=3)
    for row in range(10):
        for column in range(3):
            table.cell(row, column).text = f"row{row}-col{column}"
    data = BytesIO()
    document.save(data)

    text, _ = asyncio.run(
        extract_text(UploadFile(filename="table.docx", file=BytesIO(data.getvalue())), 20)
    )
    assert text.splitlines() == [
        f"row{row}-col{column}" for row in range(10) for column in range(3)
    ]


def test_docx_vertical_merge_is_extracted_once():
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(2, 0)).text = "Vertical once"
    table.cell(0, 1).text = "Right top"
    table.cell(1, 1).text = "Right middle"
    table.cell(2, 1).text = "Right bottom"
    data = BytesIO()
    document.save(data)

    text, _ = asyncio.run(
        extract_text(UploadFile(filename="merge.docx", file=BytesIO(data.getvalue())), 20)
    )
    assert text.splitlines() == ["Vertical once", "Right top", "Right middle", "Right bottom"]
